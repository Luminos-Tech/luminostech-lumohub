"""Fix caption for T[7] (the 1x2 signature table) and re-save."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

TARGET = r'Tino\TINO - Robot AI tiện lợi.docx'
doc = Document(TARGET)

# Find the paragraph that is the caption of T[7] (it's currently P[245] in the previous verify)
for i, p in enumerate(doc.paragraphs):
    if p.style.name == 'Caption' and '13.1' in p.text and 'ngày' in p.text:
        # Update text
        runs = p.runs
        if runs:
            runs[0].text = 'Bảng 13.1 — Cam kết và chữ ký của đại diện đội thi'
        else:
            # No run, create a new run
            r = p.add_run('Bảng 13.1 — Cam kết và chữ ký của đại diện đội thi')
        print(f'  Updated P[{i}]: {p.text!r}')
        break

doc.save(TARGET)
print('Saved fix.')
